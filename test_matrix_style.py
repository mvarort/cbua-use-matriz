import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from shutil import copyfile
from openpyxl import load_workbook
import pandas as pd
from streamlit.testing.v1 import AppTest
from matrix_style import COLORS, TARGETS, level, total_style, updated_label, display_rows
from app import excel


class MatrixStyleTests(unittest.TestCase):
    def test_all_centres_present_without_changing_counts(self):
        codes=('UX_1_1','UX_2_1','UX_2_4','UX_2_5')
        records=[{'centre':'Rector Machado','counts':dict.fromkeys(codes,2)}]
        rows=display_rows(records,codes)
        lookup={r[0]:r[1:] for r in rows}
        for centre in ['Económicas','Humanidades','Ingeniería Agronómica','Matemáticas','Bellas Artes']:
            self.assertEqual(lookup[centre],[0,0,0,0,0])
        self.assertEqual(lookup['Rector Machado'],[2,2,2,2,8])
        self.assertEqual(sum(r[-1] for r in rows),8)
        records.append({'centre':'Económicas','counts':dict.fromkeys(codes,3)})
        updated=display_rows(records,codes)
        self.assertEqual(next(r for r in updated if r[0]=='Económicas')[1:],[3,3,3,3,12])
        self.assertEqual(len([r for r in updated if r[0]=='Económicas']),1)
        book=load_workbook(BytesIO(excel(['Centro',*codes,'Total'],rows)))
        try:
            self.assertTrue(any(r[0]=='Bellas Artes' and r[-1]==0 for r in book.active.iter_rows(values_only=True)))
        finally:book.close()

    def test_all_threshold_boundaries(self):
        for centre,(first,second,third) in TARGETS.items():
            for total,expected in [(0,'Por debajo del primer objetivo'),(first-1,'Por debajo del primer objetivo'),(first,'Cumplido (con condiciones)'),(second-1,'Cumplido (con condiciones)'),(second,'Cumplido suficiente'),(third-1,'Cumplido suficiente'),(third,'Cumplido satisfactorio'),(third+1,'Cumplido satisfactorio')]:
                with self.subTest(centre=centre,total=total):
                    self.assertEqual(level(centre,total),expected)

    def test_only_total_coloured_and_unknown_neutral(self):
        row=pd.Series({'Centro':'Informática','UX_1_1':3,'UX_2_1':2,'UX_2_4':1,'UX_2_5':2,'Total':8})
        styles=total_style(row)
        self.assertEqual(styles[:-1],['']*5)
        self.assertIn(COLORS['Cumplido satisfactorio'],styles[-1])
        for name in ['Rector Machado','TOTAL','Centro FICTICIO sin objetivo']:
            row['Centro']=name
            self.assertEqual(total_style(row),['']*6)

    def test_seville_time_summer_and_winter(self):
        self.assertEqual(updated_label('2026-10-20T09:28:23.448448Z'),'20/10/2026 a las 11:28 (hora de Sevilla)')
        self.assertEqual(updated_label('2026-12-20T09:28:23Z'),'20/12/2026 a las 10:28 (hora de Sevilla)')
        with self.assertRaises(ValueError):updated_label('2026-10-20T09:28:23')

    def test_download_only_total_has_conditional_fill(self):
        columns=['Centro','UX_1_1','UX_2_1','UX_2_4','UX_2_5','Total']
        rows=[['Informática',3,2,1,2,8],['Rector Machado',1,1,1,1,4],['TOTAL',4,3,2,3,12]]
        book=load_workbook(BytesIO(excel(columns,rows)))
        try:
            sheet=book['Matriz USE']
            self.assertEqual(sheet['F2'].fill.fgColor.rgb[-6:],COLORS['Cumplido satisfactorio'][1:])
            for cell in ['A2','B2','C2','D2','E2','F3']:
                self.assertIsNone(sheet[cell].fill.patternType)
            self.assertIn('Leyenda',book.sheetnames)
            self.assertEqual(sheet['F4'].value,12)
        finally:book.close()

    def test_grouped_contract_ui_and_excel(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);(root/'data').mkdir()
            for name in ['app.py','matrix_style.py']:copyfile(Path(__file__).parent/name,root/name)
            codes=('UX_1_1','UX_2_1_2_3')
            data={'kind':'cbua_use_public_matrix_v2','updated_utc':'2026-10-20T09:28:23Z','from_inclusive':'2026-10-05','rows':[{'state':'identificado','centre':'Centro FICTICIO A','counts':{'UX_1_1':3,'UX_2_1_2_3':10}}],'totals':{'UX_1_1':3,'UX_2_1_2_3':10},'total':13}
            (root/'data/matrix.json').write_text(json.dumps(data),encoding='utf-8')
            app=AppTest.from_file(str(root/'app.py')).run()
            self.assertFalse(app.exception)
            frame=app.dataframe[0].value
            self.assertEqual(list(frame.columns),['Centro',*codes,'Total'])
            self.assertEqual(frame.iloc[-1].tolist(),['TOTAL',3,10,13])
            self.assertTrue(any('sin sumarlos' in c.value for c in app.caption))
            book=load_workbook(BytesIO(excel(list(frame.columns),frame.values.tolist())))
            try:
                self.assertEqual(book.active.max_column,4)
                self.assertTrue(book.active.auto_filter.ref.startswith('A1:D'))
                self.assertEqual(book.active.cell(book.active.max_row,4).value,13)
            finally:book.close()

    def test_synthetic_ui_and_filter(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);(root/'data').mkdir()
            for name in ['app.py','matrix_style.py']:copyfile(Path(__file__).parent/name,root/name)
            codes=('UX_1_1','UX_2_1','UX_2_4','UX_2_5')
            data={'kind':'cbua_use_public_matrix_v1','updated_utc':'2026-10-20T09:28:23Z','from_inclusive':'2026-10-05','rows':[{'state':'identificado','centre':'Centro FICTICIO A','counts':dict.fromkeys(codes,3)},{'state':'identificado','centre':'Centro FICTICIO B','counts':dict.fromkeys(codes,2)}],'totals':dict.fromkeys(codes,5),'total':20}
            (root/'data/matrix.json').write_text(json.dumps(data),encoding='utf-8')
            app=AppTest.from_file(str(root/'app.py')).run()
            self.assertFalse(app.exception)
            self.assertEqual(app.dataframe[0].value.iloc[-1,-1],20)
            self.assertTrue(any('20/10/2026 a las 11:28' in m.value for m in app.markdown))
            app.multiselect[0].set_value(['Centro FICTICIO A']).run()
            self.assertFalse(app.exception)
            self.assertEqual(app.dataframe[0].value.iloc[-1,-1],12)


if __name__=='__main__':unittest.main()
