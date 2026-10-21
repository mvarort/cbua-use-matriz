import json
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from notebook_downloads import CODES, pack, unpack, workbook



def fixture():
    books={c:{'headers':['Id','Comentario','Valor'], 'rows':[]} for c in CODES}
    for centre, number in [('Centro ficticio A',1),('Centro ficticio B',2)]:
        books['UX_2_1']['rows'].append({'centre':centre,'cells':[
            {'value':number,'type':'value','format':'0'},
            {'value':'=Texto literal\nSegunda lÃ­nea Ã±','type':'value','format':'@'},
            {'value':1.5,'type':'value','format':'0.00'}]})
    return {'kind':'cbua_use_downloads_v1','books':books,'updated_utc':'2026-10-21T12:00:00Z'}


class DownloadTests(unittest.TestCase):
    def test_public_package_roundtrip(self):
        data=fixture()
        self.assertEqual(unpack(pack(data)),data)
        with self.assertRaises(Exception):unpack(b'incompatible')

    def test_original_values_and_selection(self):
        data=fixture()
        raw,count=workbook(data['books']['UX_2_1'],'Centro ficticio A','UX_2_1')
        self.assertEqual(count,1)
        book=load_workbook(BytesIO(raw),data_only=False)
        self.assertEqual(book.active['A2'].value,1)
        self.assertEqual(book.active['B2'].value,'=Texto literal\nSegunda lÃ­nea Ã±')
        self.assertEqual(book.active['B2'].data_type,'s')
        self.assertEqual(book.active['C2'].value,1.5)
        self.assertEqual(book.active.max_row,2)
        book.close()
        self.assertEqual(workbook(data['books']['UX_1_1'],'Centro ficticio A','UX_1_1'),(None,0))

    def test_public_ui_and_matrix_independence(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);(root/'data').mkdir()
            for name in ['app.py','matrix_style.py','notebook_downloads.py']:
                shutil.copyfile(Path(__file__).parent/name,root/name)
            matrix={'kind':'cbua_use_public_matrix_v2','updated_utc':'2026-10-21T12:00:00Z',
                    'from_inclusive':'2026-10-05','rows':[],
                    'totals':{'UX_1_1':0,'UX_2_1_2_3':0},'total':0}
            (root/'data/matrix.json').write_text(json.dumps(matrix),encoding='utf-8')
            (root/'data/notebooks.json.gz').write_bytes(pack(fixture()))
            app=AppTest.from_file(str(root/'app.py')).run()
            before=app.dataframe[0].value.copy()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.text_input),0)
            self.assertEqual(len(app.selectbox),2)
            self.assertEqual(len(app.get('download_button')),0)
            app.selectbox[0].set_value('Centro ficticio A').run()
            app.selectbox[1].set_value('UX_2_1').run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.get('download_button')),1)
            self.assertTrue(before.equals(app.dataframe[0].value))
            app.selectbox[1].set_value('UX_1_1').run()
            self.assertEqual(len(app.get('download_button')),0)



if __name__=='__main__':unittest.main()
