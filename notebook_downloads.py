"""Descargas públicas de cuadernos originales filtradas por centro."""
from datetime import date, datetime, time, timedelta
from io import BytesIO
import json
import gzip
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CODES = ('UX_1_1', 'UX_2_1', 'UX_2_4', 'UX_2_5')


def pack(package):
    raw = json.dumps(package, ensure_ascii=False, allow_nan=False).encode('utf-8')
    return gzip.compress(raw, mtime=0)


def unpack(blob):
    result = json.loads(gzip.decompress(blob))
    if result['kind'] != 'cbua_use_downloads_v1' or set(result['books']) != set(CODES):
        raise ValueError('PAQUETE_INCOMPATIBLE')
    return result



def workbook(sheet, centre, code):
    rows = [r['cells'] for r in sheet['rows'] if r['centre'] == centre]
    if not rows:
        return None, 0
    book = Workbook()
    ws = book.active
    ws.title = code
    ws.append(sheet['headers'])
    for row in rows:
        ws.append([None] * len(sheet['headers']))
        for index, original in enumerate(row, 1):
            cell = ws.cell(ws.max_row, index)
            value = original['value']
            if original['type'] == 'datetime':
                value = datetime.fromisoformat(value)
            elif original['type'] == 'date':
                value = date.fromisoformat(value)
            elif original['type'] == 'time':
                value = time.fromisoformat(value)
            elif original['type'] == 'timedelta':
                value = timedelta(seconds=value)
            cell.value = value
            # Los textos que empiezan por '=' se exportan literalmente, nunca como fórmulas.
            if isinstance(value, str):
                cell.data_type = 's'
            cell.number_format = original['format'] or 'General'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for cell in ws[1]:
        cell.data_type = 's'
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='9E354B')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for index in range(1, ws.max_column+1):
        ws.column_dimensions[get_column_letter(index)].width = 35
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue(), len(rows)


def sidebar_downloads(root):
    import streamlit as st
    st.sidebar.markdown('### Descargar un cuaderno')
    target = Path(root)/'data/notebooks.json.gz'
    if not target.is_file():
        st.sidebar.info('Descargas pendientes de preparar.')
        return
    try:
        package = unpack(target.read_bytes())
    except Exception:
        st.sidebar.error('El archivo de descargas no está disponible o es incompatible.')
        return
    centres = sorted({r['centre'] for sheet in package['books'].values() for r in sheet['rows']})
    centre = st.sidebar.selectbox('Centro para la descarga', centres, index=None,
                                  placeholder='Selecciona un centro')
    code = st.sidebar.selectbox('Cuaderno UX', CODES, index=None,
                                placeholder='Selecciona un cuaderno')
    if centre is None or code is None:
        return
    content, count = workbook(package['books'][code], centre, code)
    st.sidebar.caption('Datos de descarga: '+package['updated_utc'])
    if not count:
        st.sidebar.info('Sin registros para esta selección.')
        return
    st.sidebar.caption(f'{count} registros · {centre} · {code}')
    safe = ''.join(c if c.isalnum() else '_' for c in centre)
    st.sidebar.download_button('Descargar cuaderno en Excel', content,
        file_name=f'{safe}_{code}.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        key=f'download_{safe}_{code}')
