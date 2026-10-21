"""Matriz pública: exclusivamente recuentos agregados por centro y cuaderno."""
from io import BytesIO
import json
from pathlib import Path
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from notebook_downloads import sidebar_downloads
from matrix_style import COLORS, TARGETS, level, total_style, updated_label, display_rows, cell_levels, COLUMN_LABELS

LEGACY_CODES=('UX_1_1','UX_2_1','UX_2_4','UX_2_5')
CODES=('UX_1_1','UX_2_1_2_3')

def package_codes(data):
    return CODES if data['kind']=='cbua_use_public_matrix_v2' else LEGACY_CODES


def validate(data):
    if set(data)!={'kind','updated_utc','from_inclusive','rows','totals','total'} or data['kind'] not in {'cbua_use_public_matrix_v1','cbua_use_public_matrix_v2'} or data['from_inclusive']!='2026-10-05':
        raise ValueError('Contrato incompatible')
    codes=package_codes(data)
    if set(data['totals'])!=set(codes):
        raise ValueError('Cuadernos incompatibles')
    seen=set()
    for row in data['rows']:
        if set(row)!={'state','centre','counts'} or row['state']!='identificado' or set(row['counts'])!=set(codes):
            raise ValueError('Fila incompatible')
        if row['state']=='identificado' and (not isinstance(row['centre'],str) or not row['centre'].strip()):
            raise ValueError('Centro incompatible')
        if row['state']!='identificado' and row['centre'] is not None:
            raise ValueError('Centro incompatible')
        key=(row['state'],row['centre'])
        if key in seen or any(type(n) is not int or n<0 for n in row['counts'].values()):
            raise ValueError('Recuento incompatible')
        seen.add(key)
    if any(type(data['totals'][c]) is not int or data['totals'][c]!=sum(r['counts'][c] for r in data['rows']) for c in codes):
        raise ValueError('Totales contradictorios')
    if type(data['total']) is not int or data['total']!=sum(data['totals'].values()):
        raise ValueError('Total contradictorio')


def excel(columns,rows):
    book=Workbook();sheet=book.active;sheet.title='Matriz USE'
    headers=[COLUMN_LABELS.get(c,c) if 'UX_2_1_2_3' in columns else c for c in columns]
    for i,row in enumerate([headers,*rows],1):
        for j,value in enumerate(row,1):
            cell=sheet.cell(i,j,value)
            if isinstance(value,str):cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if i==1 or i==len(rows)+1:
                cell.font=Font(bold=True,color='FFFFFF' if i==1 else '6E2332')
                if i==1:cell.fill=PatternFill('solid',fgColor='9E354B')
    for number, row in enumerate(rows, 2):
        states=cell_levels(dict(zip(columns,row)))
        for column,state in states.items():
            sheet.cell(number,columns.index(column)+1).fill=PatternFill('solid',fgColor=COLORS[state].lstrip('#'))
    legend=book.create_sheet('Leyenda')
    legend.append(['Aplicación','Total: UX 1_1 + máximo(UX 2_1, UX 2_4, UX 2_5)' if 'UX_2_1_2_3' in columns else 'Total de los cuatro cuadernos'])
    for state,color in COLORS.items():
        legend.append([state]);legend.cell(legend.max_row,1).fill=PatternFill('solid',fgColor=color.lstrip('#'))
    legend.append(['Sin relleno','Centro sin objetivo definido o total general'])
    legend.append(['Centro','Con condiciones','Suficiente','Satisfactorio'])
    for centre,limits in TARGETS.items():legend.append([centre,*limits])
    legend.column_dimensions['A'].width=44
    legend.column_dimensions['B'].width=55
    sheet.freeze_panes='B2'
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{len(rows)}'
    for number in range(2,len(columns)+1):sheet.column_dimensions[get_column_letter(number)].width=22
    sheet.column_dimensions['A'].width=48
    result=BytesIO();book.save(result);book.close()
    return result.getvalue()


def main():
    st.set_page_config(page_title='Matriz USE | CBUA',page_icon='📚',layout='wide')
    st.markdown('<style>[data-testid="stAppViewContainer"]{background:#FCFAFA;}[data-testid="stSidebar"]{background:#F6EDEF;}h1,h2,h3{color:#8C293F;}[data-testid="stSidebar"]{border-right:1px solid #EAD9DE;}button[kind="secondary"]{border-color:#D7B8C0;}.block-container{max-width:1250px;padding-top:3rem;}</style>',unsafe_allow_html=True)
    st.sidebar.markdown('### Biblioteca de la Universidad de Sevilla')
    st.title('Matriz de observaciones USE')
    st.caption('Resultados provisionales — estudios UX en curso')
    try:
        data=json.loads((Path(__file__).parent/'data/matrix.json').read_text(encoding='utf-8'))
        validate(data)
    except Exception:
        st.error('La matriz no está disponible o sus totales son incompatibles.')
        st.stop()
    st.markdown('**Última actualización de los datos: '+updated_label(data['updated_utc'])+'**')
    st.caption('Observaciones desde el 05/10/2026 · Actualización diaria programada a las 09:00 (Sevilla)')
    codes=package_codes(data)
    sidebar_downloads(Path(__file__).parent)
    source=display_rows(data['rows'], codes)
    rows=list(source)
    rows.append(['TOTAL',*[sum(r[i] for r in rows) for i in range(1,len(codes)+2)]])
    columns=['Centro',*codes,'Total']
    config={c:st.column_config.NumberColumn(COLUMN_LABELS.get(c,c)) for c in codes} if data['kind']=='cbua_use_public_matrix_v2' else {}
    st.dataframe(pd.DataFrame(rows,columns=columns).style.apply(total_style,axis=1),hide_index=True,column_config=config,width='stretch',height=min(900,max(180,40+35*len(rows))))


if __name__=='__main__':
    main()
